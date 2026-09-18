from __future__ import annotations

import html
import re
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE_DOCX = ROOT / "Huong-dan-cau-hinh-he-thong-VIN-HRM.docx"
OUT_HTML = ROOT / "Huong-dan-cau-hinh-he-thong-VIN-HRM.html"

FIGURES = [
    (
        "images/01-doanh-nghiep-trai.png",
        "images/01-doanh-nghiep-phai.png",
        "Thẻ Doanh nghiệp của màn hình Cấu hình hệ thống VIN-HRM",
    ),
    (
        "images/02-chi-nhanh-trai.png",
        "images/02-chi-nhanh-phai.png",
        "Thẻ Chi nhánh với chi nhánh Trụ sở Hà Nội đang được chọn",
    ),
    (
        "images/03-chon-chi-nhanh-trai.png",
        "images/03-chon-chi-nhanh-phai.png",
        "Danh sách chọn chi nhánh đang hoạt động",
    ),
    (
        "images/04-lich-su-trai.png",
        "images/04-lich-su-phai.png",
        "Lịch sử cấu hình theo doanh nghiệp",
    ),
]


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent")

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def has_page_break(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath('.//w:br[@w:type="page"]'))


def has_numbering(paragraph: Paragraph) -> bool:
    ppr = paragraph._p.pPr
    return bool(ppr is not None and ppr.numPr is not None)


def has_drawing(table: Table) -> bool:
    return bool(table._tbl.xpath(".//w:drawing"))


def paragraph_html(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    safe = html.escape(text)
    if "02. Cấu hình v1.0" in text:
        safe = safe.replace(
            "02. Cấu hình v1.0",
            '<a href="https://docs.google.com/document/d/11dW_3cQNBK_E86RgZMj4v8VdcBwwzfNQYNPvenrjcPU/edit">02. Cấu hình v1.0</a>',
        )
    return safe


def callout_class(text: str) -> str:
    warning_starts = (
        "LƯU Ý",
        "Không có màn hình sửa mặc định",
        "Kiểm tra quyền",
        "Ràng buộc ngày trong tháng",
        "Tải lại",
        "Kiểm tra tên chi nhánh",
        "Khôi phục kế thừa",
        "Nguyên tắc",
    )
    success_starts = ("Ví dụ thực tế", "Bảo vệ thông tin nhạy cảm")
    if text.startswith(warning_starts):
        return "warning"
    if text.startswith(success_starts):
        return "success"
    return "note"


def render_callout(table: Table) -> str:
    text = " ".join(p.text.strip() for p in table.cell(0, 0).paragraphs if p.text.strip())
    safe = html.escape(text)
    if ":" in safe:
        label, rest = safe.split(":", 1)
        safe = f"<strong>{label}:</strong>{rest}"
    kind = callout_class(text)
    return f'<table class="callout {kind}"><tr><td>{safe}</td></tr></table>'


def render_figure(index: int) -> str:
    left, right, alt = FIGURES[index]
    alt_safe = html.escape(alt, quote=True)
    return (
        '<table class="figure-pair" role="presentation"><tr>'
        f'<td><img src="{left}" alt="{alt_safe} - nửa trái"></td>'
        f'<td><img src="{right}" alt="{alt_safe} - nửa phải"></td>'
        "</tr></table>"
    )


def render_data_table(table: Table, cover: bool = False) -> str:
    klass = "meta-table" if cover else "data-table"
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for cell_index, cell in enumerate(row.cells):
            tag = "th" if (cover and cell_index == 0) or (not cover and row_index == 0) else "td"
            cell_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append(f"<{tag}>{html.escape(cell_text)}</{tag}>")
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return f'<table class="{klass}">' + "".join(rows) + "</table>"


def build_html() -> None:
    document = Document(SOURCE_DOCX)
    blocks = list(iter_block_items(document))
    output: list[str] = []
    list_items: list[str] = []
    list_start = 1
    list_counter = 1
    figure_index = 0
    cover = True
    toc_mode = False

    def flush_list() -> None:
        nonlocal list_items, list_start
        if not list_items:
            return
        items = "".join(f"<li>{item}</li>" for item in list_items)
        output.append(f'<ol class="steps" start="{list_start}">{items}</ol>')
        list_items = []

    output.append('<section class="cover">')

    for block in blocks:
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if has_page_break(block):
                flush_list()
                if cover:
                    output.append("</section>")
                    cover = False
                output.append('<div class="hard-break"></div>')
                continue
            if not text:
                continue
            if has_numbering(block):
                if not list_items:
                    list_start = list_counter
                list_items.append(paragraph_html(block))
                list_counter += 1
                continue

            flush_list()
            style = block.style.name if block.style else "Normal"
            safe = paragraph_html(block)

            if text == "Mục lục":
                toc_mode = True
            elif text == "1. Mục đích và phạm vi":
                toc_mode = False

            if cover:
                if text == "HƯỚNG DẪN SỬ DỤNG":
                    output.append(f'<p class="cover-kicker">{safe}</p>')
                elif text == "CẤU HÌNH HỆ THỐNG":
                    output.append(f'<h1 class="cover-title">{safe}</h1>')
                elif text.startswith("Cấu hình mặc định -"):
                    output.append(f'<p class="cover-subtitle">{safe}</p>')
                elif text.startswith("Tài liệu dành cho"):
                    output.append(f'<p class="cover-audience">{safe}</p>')
                elif text.startswith("VIN-HRM - Quản trị"):
                    output.append(f'<p class="cover-brand">{safe}</p>')
                else:
                    output.append(f"<p>{safe}</p>")
            elif style.startswith("Heading 1"):
                output.append(f"<h1>{safe}</h1>")
            elif style.startswith("Heading 2"):
                output.append(f"<h2>{safe}</h2>")
            elif style.startswith("Heading 3"):
                output.append(f"<h3>{safe}</h3>")
            elif style == "Caption":
                output.append(f"<figcaption>{safe}</figcaption>")
            elif text == "Ảnh chụp giao diện thực tế":
                output.append(f'<p class="figure-label">{safe}</p>')
            elif toc_mode and re.match(r"^\d+\.\s", text):
                output.append(f'<p class="toc-item">{safe}</p>')
            else:
                output.append(f"<p>{safe}</p>")

        elif isinstance(block, Table):
            flush_list()
            if has_drawing(block):
                output.append(render_figure(figure_index))
                figure_index += 1
            elif len(block.rows) == 1 and len(block.columns) == 1:
                output.append(render_callout(block))
            else:
                output.append(render_data_table(block, cover=cover))

    flush_list()
    if cover:
        output.append("</section>")

    body = "\n".join(output)
    page_title = "Hướng dẫn cấu hình hệ thống VIN-HRM"
    html_doc = f"""<!doctype html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="author" content="VIN-HRM">
  <title>{page_title}</title>
  <style>
    @page {{ size: Letter portrait; margin: 0.68in 0.72in 0.7in; }}
    * {{ box-sizing: border-box; }}
    html {{ background: #e8edf3; }}
    body {{ margin: 0; color: #243142; font-family: Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.34; }}
    .document {{ width: 8.5in; margin: 24px auto; padding: 0.68in 0.72in 0.7in; background: white; box-shadow: 0 8px 28px rgba(15, 35, 60, .15); }}
    p {{ margin: 0 0 7pt; orphans: 3; widows: 3; }}
    h1, h2, h3 {{ color: #173b67; break-after: avoid-page; page-break-after: avoid; }}
    h1 {{ margin: 18pt 0 10pt; font-size: 17pt; line-height: 1.15; }}
    h2 {{ margin: 14pt 0 7pt; font-size: 13pt; line-height: 1.2; }}
    h3 {{ margin: 10pt 0 5pt; font-size: 12pt; }}
    a {{ color: #2f78bc; text-decoration: underline; }}
    .cover {{ min-height: 9.2in; text-align: center; padding-top: 0.75in; page-break-after: always; break-after: page; }}
    .cover-kicker {{ margin-top: 0.2in; color: #2f78bc; font-size: 11pt; font-weight: 700; letter-spacing: .03em; }}
    .cover-title {{ margin: 24pt 0 14pt; color: #173b67; font-size: 28pt; line-height: 1.1; }}
    .cover-subtitle {{ margin-bottom: 28pt; color: #2b5b8b; font-size: 14pt; }}
    .cover-audience {{ margin: 0 0 36pt; color: #6c7580; font-style: italic; }}
    .cover-brand {{ margin-top: 34pt; color: #2f78bc; font-weight: 700; }}
    .hard-break {{ page-break-before: always; break-before: page; height: 0; }}
    .toc-item {{ margin: 0 0 6pt 12pt; color: #244f7d; }}
    ol.steps {{ margin: 4pt 0 9pt 19pt; padding-left: 15pt; }}
    ol.steps li {{ margin: 0 0 6pt 2pt; padding-left: 3pt; break-inside: avoid; page-break-inside: avoid; }}
    table {{ width: 100%; border-collapse: collapse; margin: 8pt 0 12pt; break-inside: auto; page-break-inside: auto; }}
    tr {{ break-inside: avoid; page-break-inside: avoid; }}
    th, td {{ border: 1px solid #c9d5e4; padding: 6pt 7pt; vertical-align: middle; text-align: left; }}
    th {{ background: #e4ecf5; color: #173b67; font-weight: 700; }}
    .meta-table {{ margin: 18pt 0 10pt; }}
    .meta-table th {{ width: 25%; }}
    .figure-pair {{ table-layout: fixed; margin: 7pt 0 4pt; break-inside: avoid; page-break-inside: avoid; }}
    .figure-pair td {{ width: 50%; padding: 0; border: 0; vertical-align: top; }}
    .figure-pair img {{ display: block; width: 100%; height: auto; }}
    .figure-label {{ margin: 7pt 0 4pt; color: #2f78bc; font-size: 9pt; font-weight: 700; break-after: avoid-page; page-break-after: avoid; }}
    figcaption {{ margin: 0 0 8pt; text-align: center; font-size: 9pt; font-style: italic; font-weight: 700; color: #384658; break-after: avoid; }}
    .callout {{ margin: 9pt 0 12pt; break-inside: avoid; page-break-inside: avoid; }}
    .callout td {{ padding: 8pt 9pt; border-width: 1.2pt; }}
    .callout.note td {{ background: #f2f5f8; border-color: #c3d0df; }}
    .callout.warning td {{ background: #fff3cd; border-color: #d8a817; }}
    .callout.success td {{ background: #eaf6ec; border-color: #54a46b; }}
    .callout strong {{ color: #173b67; }}
    @media print {{
      html, body {{ background: white; }}
      .document {{ width: auto; margin: 0; padding: 0; box-shadow: none; }}
    }}
  </style>
</head>
<body>
  <main class="document Section1">
{body}
  </main>
</body>
</html>
"""
    OUT_HTML.write_text(html_doc, encoding="utf-8")
    print(OUT_HTML)
    print(f"figures={figure_index}; list_items={list_counter - 1}")


if __name__ == "__main__":
    build_html()
