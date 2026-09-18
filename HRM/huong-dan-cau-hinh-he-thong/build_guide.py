from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\Works\HRM\output\huong-dan-cau-hinh-he-thong")
IMG = ROOT / "images"
OUT = ROOT / "Huong-dan-cau-hinh-he-thong-VIN-HRM.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
TEXT = "202124"
MUTED = "5F6368"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "D7DEE8"
CAUTION = "FFF4CE"
CAUTION_BORDER = "D6B656"
SUCCESS = "E8F5E9"
SUCCESS_BORDER = "7CB38B"
WHITE = "FFFFFF"

CONTENT_DXA = 9360
FIGURE_HALF_DXA = 4680
TABLE_INDENT_DXA = 120


def set_run_font(run, name="Calibri", size=None, color=TEXT, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(idx, len(widths_dxa) - 1)])


def set_table_borders(table, color=BORDER, size="4", no_borders=False):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil" if no_borders else "single")
        if not no_borders:
            node.set(qn("w:sz"), size)
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_paragraph_border(paragraph, color=BORDER, size="6", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(x.get(qn("w:abstractNumId")))
        for x in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id_node = OxmlElement("w:abstractNumId")
    abstract_id_node.set(qn("w:val"), str(abstract_id))
    num.append(abstract_id_node)
    numbering.append(num)
    return num_id, abstract_id


def new_numbering_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_bullet(doc, text, bullet_num_id):
    p = doc.add_paragraph(style="Normal")
    apply_num(p, bullet_num_id)
    p.add_run(text)
    return p


def add_step(doc, text, num_id):
    p = doc.add_paragraph(style="Normal")
    apply_num(p, num_id)
    p.add_run(text)
    return p


def add_body(doc, text, bold_lead=None, keep_with_next=False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_callout(doc, label, text, kind="note"):
    fill = LIGHT_GRAY
    border = BORDER
    if kind == "warning":
        fill, border = CAUTION, CAUTION_BORDER
    elif kind == "success":
        fill, border = SUCCESS, SUCCESS_BORDER
    table = doc.add_table(rows=1, cols=1)
    prevent_row_split(table.rows[0])
    set_table_geometry(table, [CONTENT_DXA], indent_dxa=TABLE_INDENT_DXA)
    set_table_borders(table, color=border, size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=NAVY)
    r2 = p.add_run(text)
    set_run_font(r2, color=TEXT)
    return table


def add_table(doc, headers, rows, widths_dxa, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=TEXT)
    return table


def add_figure_pair(doc, left_name, right_name, caption, alt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("Ảnh chụp giao diện thực tế")
    set_run_font(r, size=9, bold=True, color=BLUE)

    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [FIGURE_HALF_DXA, FIGURE_HALF_DXA], indent_dxa=0)
    set_table_borders(table, no_borders=True)
    row = table.rows[0]
    prevent_row_split(row)
    for idx, name in enumerate((left_name, right_name)):
        cell = row.cells[idx]
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        p_img = cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(0)
        run = p_img.add_run()
        inline = run.add_picture(str(IMG / name), width=Inches(3.25))
        inline._inline.docPr.set("descr", f"{alt} - nửa {'trái' if idx == 0 else 'phải'}")

    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    cap.add_run(caption)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.keep_together = True


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("VIN-HRM  |  HƯỚNG DẪN CẤU HÌNH HỆ THỐNG")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    footer = section.footer
    p0 = footer.paragraphs[0]
    p0.text = ""
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [7200, 2160], indent_dxa=0)
    set_table_borders(table, no_borders=True)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r1 = left.add_run("Tài liệu hướng dẫn sử dụng - Phiên bản 1.0")
    set_run_font(r1, size=8.5, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    r2 = right.add_run("Trang ")
    set_run_font(r2, size=8.5, color=MUTED)
    add_field(right, "PAGE")
    r3 = right.add_run(" / ")
    set_run_font(r3, size=8.5, color=MUTED)
    add_field(right, "NUMPAGES")

    first_footer = section.first_page_footer
    first_footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_footer.paragraphs[0].paragraph_format.space_after = Pt(0)
    fr = first_footer.paragraphs[0].add_run("VIN-HRM - Tài liệu hướng dẫn sử dụng")
    set_run_font(fr, size=8.5, color=MUTED)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(68)
    p.paragraph_format.space_after = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HƯỚNG DẪN SỬ DỤNG")
    set_run_font(r, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("CẤU HÌNH HỆ THỐNG")
    set_run_font(r, size=30, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("Cấu hình mặc định - Cấu hình theo doanh nghiệp - Cấu hình theo chi nhánh")
    set_run_font(r, size=14, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(46)
    r = p.add_run("Tài liệu dành cho người quản trị hệ thống VIN-HRM")
    set_run_font(r, size=11, color=MUTED, italic=True)

    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [2400, 6960], indent_dxa=120)
    set_table_borders(table, color=BORDER, size="4")
    data = [
        ("Phiên bản", "1.0"),
        ("Ngày đối chiếu", "13/08/2026"),
        ("Môi trường kiểm tra", "Mã nguồn cục bộ và dữ liệu nghiệp vụ trên hệ thống VIN-HRM"),
    ]
    for row, (label, value) in zip(table.rows, data):
        prevent_row_split(row)
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell, top=100, start=140, bottom=100, end=140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        rl = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(rl, size=10, bold=True, color=NAVY)
        rv = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(rv, size=10, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VIN-HRM - Quản trị nhân sự toàn diện")
    set_run_font(r, size=10.5, color=BLUE, bold=True)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_page(section)
    configure_header_footer(section)
    props = doc.core_properties
    props.title = "Hướng dẫn cấu hình hệ thống VIN-HRM"
    props.subject = "Cấu hình mặc định, cấu hình theo doanh nghiệp và cấu hình theo chi nhánh"
    props.author = "VIN-HRM"
    props.keywords = "VIN-HRM, cấu hình hệ thống, doanh nghiệp, chi nhánh"
    props.comments = "Đối chiếu tài liệu đặc tả v1.0 và giao diện phần mềm thực tế ngày 13/08/2026."

    bullet_num_id, _ = add_numbering_definition(doc, "bullet")
    _, decimal_abstract_id = add_numbering_definition(doc, "decimal")

    add_cover(doc)
    doc.add_page_break()

    add_heading(doc, "Mục lục", 1)
    toc_items = [
        "1. Mục đích và phạm vi",
        "2. Cách hệ thống xác định giá trị cấu hình",
        "3. Truy cập màn hình Cấu hình hệ thống",
        "4. Cấu hình mặc định hệ thống",
        "5. Cấu hình theo doanh nghiệp",
        "6. Cấu hình theo chi nhánh",
        "7. Kiểm tra lịch sử thay đổi",
        "8. Tình huống thường gặp và cách xử lý",
        "9. Checklist vận hành an toàn",
        "10. Phụ lục",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(item)
        set_run_font(r, size=11, color=NAVY)
    add_callout(
        doc,
        "Cách đọc tài liệu",
        "Thực hiện lần lượt từ phần 2 đến phần 7 khi mới làm quen. Nếu đã biết đường dẫn màn hình, có thể đi thẳng tới phần cấu hình theo doanh nghiệp hoặc chi nhánh.",
    )

    add_heading(doc, "1. Mục đích và phạm vi", 1)
    add_body(
        doc,
        "Tài liệu này hướng dẫn người quản trị sử dụng màn hình Cấu hình hệ thống của VIN-HRM để hiểu giá trị mặc định, thiết lập giá trị áp dụng cho toàn doanh nghiệp, thiết lập giá trị riêng cho từng chi nhánh và kiểm tra lịch sử thay đổi.",
    )
    add_body(
        doc,
        "Nội dung được đối chiếu từ tài liệu đặc tả nghiệp vụ '02. Cấu hình v1.0' và từ phần mềm thực tế chạy bằng mã nguồn tại thời điểm ngày 13/08/2026. Tài khoản dùng để kiểm tra giao diện là tài khoản quản trị hrm_admin; mật khẩu không được ghi lại trong tài liệu.",
    )
    add_body(doc, "Tài liệu tập trung vào các thao tác người dùng, không yêu cầu người đọc biết lập trình hoặc cấu trúc cơ sở dữ liệu.")

    add_heading(doc, "1.1. Đối tượng sử dụng", 2)
    for text in (
        "Quản trị viên hệ thống chịu trách nhiệm cấu hình các quy tắc dùng chung.",
        "Người phụ trách nhân sự cần hiểu giá trị đang được áp dụng để kiểm tra nghiệp vụ.",
        "Người hỗ trợ vận hành cần tra cứu lịch sử khi một quy tắc thay đổi ngoài dự kiến.",
    ):
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "1.2. Điều quan trọng cần nhớ", 2)
    add_callout(
        doc,
        "LƯU Ý",
        "Phần mềm hiện có ba thẻ Doanh nghiệp, Chi nhánh và Lịch sử. Không có thẻ riêng mang tên 'Cấu hình mặc định'. Giá trị mặc định được hệ thống định nghĩa sẵn và được dùng tự động khi chưa có giá trị ghi đè.",
        kind="warning",
    )

    add_heading(doc, "2. Cách hệ thống xác định giá trị cấu hình", 1)
    add_body(
        doc,
        "Mỗi cấu hình được nhận diện bằng một mã khóa duy nhất. Khi cần sử dụng cấu hình, hệ thống tìm giá trị theo thứ tự ưu tiên từ phạm vi hẹp nhất đến phạm vi rộng nhất.",
    )
    add_table(
        doc,
        ["Ưu tiên", "Phạm vi", "Khi nào được sử dụng"],
        [
            ("1 - Cao nhất", "Chi nhánh", "Dùng khi chi nhánh đang làm việc có giá trị cấu hình riêng."),
            ("2", "Doanh nghiệp", "Dùng khi chi nhánh chưa có giá trị riêng nhưng doanh nghiệp đã cấu hình."),
            ("3 - Nền", "Mặc định hệ thống", "Dùng khi cả chi nhánh và doanh nghiệp đều chưa có giá trị ghi đè."),
        ],
        [1500, 2200, 5660],
    )
    add_callout(
        doc,
        "Ví dụ thực tế",
        "Trong ảnh chụp, doanh nghiệp đang có 'Ngày tự động xếp lịch trong tuần' là Thứ 2, còn chi nhánh Trụ sở Hà Nội hiển thị Thứ 4. Khi nghiệp vụ chạy cho Trụ sở Hà Nội, hệ thống dùng Thứ 4 vì cấu hình chi nhánh có ưu tiên cao hơn.",
        kind="success",
    )

    add_heading(doc, "2.1. Cấu hình chỉ áp dụng ở cấp doanh nghiệp", 2)
    add_body(
        doc,
        "Một số cấu hình không được phép thay đổi theo chi nhánh. Trên phần mềm thực tế, các thông tin kết nối bảo hiểm xã hội và dung lượng tối đa mỗi tệp trong Chat nội bộ chỉ xuất hiện ở thẻ Doanh nghiệp. Đây là hành vi đúng; không phải dữ liệu chi nhánh bị thiếu.",
    )

    add_heading(doc, "3. Truy cập màn hình Cấu hình hệ thống", 1)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Mở VIN-HRM và đăng nhập bằng tài khoản được cấp quyền quản trị.",
        "Trong thanh điều hướng bên trái, mở nhóm Hệ thống.",
        "Chọn Cấu hình. Tiêu đề ở đầu trang phải hiển thị CẤU HÌNH HỆ THỐNG.",
        "Kiểm tra ba thẻ Doanh nghiệp, Chi nhánh và Lịch sử trước khi thao tác.",
    ):
        add_step(doc, text, steps)
    add_callout(
        doc,
        "Kiểm tra quyền",
        "Nếu không nhìn thấy mục Cấu hình hoặc chỉ có thể xem mà không thể lưu, hãy liên hệ người quản trị phân quyền. Không dùng tài khoản của người khác để thay đổi cấu hình.",
        kind="warning",
    )

    doc.add_page_break()
    add_heading(doc, "4. Cấu hình mặc định hệ thống", 1)
    add_heading(doc, "4.1. Cấu hình mặc định là gì?", 2)
    add_body(
        doc,
        "Cấu hình mặc định là tập hợp các giá trị nền do hệ thống định nghĩa sẵn theo nghiệp vụ. Danh sách này được khai báo trong mã nguồn và đưa vào cơ sở dữ liệu trong quá trình khởi tạo hoặc nâng cấp hệ thống.",
    )
    add_body(
        doc,
        "Người dùng không tạo mới, sửa hoặc xóa cấu hình mặc định trực tiếp trên giao diện. Khi doanh nghiệp chưa thiết lập một cấu hình, thẻ Doanh nghiệp hiển thị giá trị hiệu lực được lấy từ mặc định hệ thống. Tương tự, khi chi nhánh chưa có giá trị riêng, hệ thống tiếp tục tìm giá trị của doanh nghiệp rồi mới dùng mặc định.",
    )
    add_callout(
        doc,
        "Không có màn hình sửa mặc định",
        "Nếu cần thay đổi giá trị nền cho toàn bộ khách hàng hoặc toàn bộ hệ thống, yêu cầu này phải được xử lý trong quy trình phát triển và triển khai phần mềm. Không tự sửa trực tiếp dữ liệu nền.",
        kind="warning",
    )

    add_heading(doc, "4.2. Các dạng giá trị hiển thị", 2)
    add_table(
        doc,
        ["Dạng cấu hình", "Cách hiển thị trên giao diện", "Cách sử dụng"],
        [
            ("Danh sách chọn", "Ô có mũi tên xổ xuống", "Chọn một giá trị hợp lệ do hệ thống cung cấp."),
            ("Bật / tắt", "Công tắc", "Bật hoặc tắt quy tắc. Công tắc màu xanh thường biểu thị đang bật."),
            ("Số", "Ô nhập số", "Chỉ nhập số trong phạm vi nghiệp vụ cho phép."),
            ("Thời gian", "Ô chọn giờ", "Chọn hoặc nhập thời gian theo định dạng của trình duyệt."),
            ("Văn bản", "Ô nhập chữ", "Nhập nội dung theo yêu cầu của cấu hình."),
            ("Mật khẩu", "Ô nhập được che", "Nhập bí mật; có thể dùng biểu tượng con mắt để kiểm tra tạm thời."),
        ],
        [1900, 2600, 4860],
    )

    add_heading(doc, "4.3. Một số ví dụ cấu hình nền", 2)
    add_body(
        doc,
        "Tùy phiên bản, danh sách có thể thay đổi. Các nhóm quan sát được trên phần mềm gồm Lịch làm việc, Chấm công, Chat nội bộ, Bảo hiểm xã hội và Hệ thống. Tài liệu đặc tả cũng mô tả các mã như chu kỳ xếp lịch, định dạng ngày, khoảng thời gian check-in sớm và mức giảm trừ khi tính thuế.",
    )
    add_body(
        doc,
        "Các ví dụ thường gặp gồm chu kỳ tự động xếp lịch làm việc; ngày hoặc giờ tự động xếp lịch; quy tắc ghi nhận check-out trễ; định dạng ngày và ngày giờ; cùng các công tắc cho Chat nội bộ, tính lương hoặc OKR.",
    )
    doc.add_page_break()
    add_heading(doc, "5. Cấu hình theo doanh nghiệp", 1)
    add_body(
        doc,
        "Cấu hình theo doanh nghiệp áp dụng cho toàn bộ doanh nghiệp và cho tất cả chi nhánh chưa có thiết lập riêng. Đây là cấp nên dùng cho các quy tắc chung như định dạng ngày, nguyên tắc tính lương, chính sách Chat nội bộ hoặc lịch tự động dùng chung. Giá trị trên thẻ là giá trị hiệu lực; giao diện hiện chưa gắn nhãn để cho biết từng dòng đến từ mặc định hệ thống hay một lần ghi đè trước đó.",
    )
    add_figure_pair(
        doc,
        "01-doanh-nghiep-trai.png",
        "01-doanh-nghiep-phai.png",
        "Hình 1. Thẻ Doanh nghiệp trên giao diện Cấu hình hệ thống (khung giao diện rộng 1920 px).",
        "Thẻ Doanh nghiệp của màn hình Cấu hình hệ thống VIN-HRM",
    )
    add_callout(
        doc,
        "Các vị trí cần chú ý",
        "Phía trên là ba thẻ phạm vi. Thanh công cụ có ô Tìm cấu hình, danh sách nhóm, nút Xóa bộ lọc, Tải lại và Lưu thay đổi. Phần dưới được chia thành các nhóm cấu hình để dễ theo dõi.",
    )

    add_heading(doc, "5.1. Tìm và lọc cấu hình", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Chọn thẻ Doanh nghiệp.",
        "Nếu biết tên hoặc mã cấu hình, nhập từ khóa vào ô Tìm cấu hình. Có thể tìm theo mô tả, mã khóa hoặc loại dữ liệu.",
        "Nếu muốn xem theo nhóm, mở danh sách đang hiển thị Tất cả rồi chọn Workflow, Lịch làm việc, Chấm công, Chat nội bộ, Bảo hiểm xã hội hoặc Hệ thống.",
        "Chọn Xóa bộ lọc để xóa từ khóa và trở lại nhóm Tất cả.",
    ):
        add_step(doc, text, steps)
    add_callout(
        doc,
        "Xóa bộ lọc không xóa dữ liệu",
        "Nút Xóa bộ lọc chỉ đưa điều kiện tìm kiếm về ban đầu. Nút này không hoàn tác các giá trị đã nhập và không xóa cấu hình đã lưu.",
    )

    add_heading(doc, "5.2. Chỉnh sửa giá trị", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Xác định đúng dòng cấu hình qua cột Tên cấu hình.",
        "Thay đổi giá trị ở cột Giá trị bằng danh sách chọn, công tắc, ô số, ô thời gian hoặc ô văn bản tương ứng.",
        "Rà soát các trường liên quan. Ví dụ, khi chu kỳ xếp lịch là Hằng tuần, hệ thống hiển thị ngày trong tuần; khi chọn Hằng tháng, phải chọn ngày trong tháng.",
        "Kiểm tra lại ảnh hưởng đối với toàn bộ chi nhánh trước khi lưu.",
    ):
        add_step(doc, text, steps)
    add_callout(
        doc,
        "Ràng buộc ngày trong tháng",
        "Nếu chọn chu kỳ Hằng tháng, ngày tự động xếp lịch phải từ 1 đến 31. Khi nhập sai, hệ thống hiển thị: 'Ngày tự động xếp lịch trong tháng chỉ được nhập giá trị từ 1 đến 31.'",
        kind="warning",
    )

    add_heading(doc, "5.3. Lưu cấu hình doanh nghiệp", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Sau khi kiểm tra, chọn Lưu thay đổi ở góc phải thanh công cụ.",
        "Chờ hệ thống xử lý; không nhấn nhiều lần liên tiếp.",
        "Khi thành công, hệ thống hiển thị thông báo 'Lưu cấu hình doanh nghiệp thành công'.",
        "Mở thẻ Lịch sử để xác nhận giá trị cũ và giá trị mới nếu thay đổi quan trọng.",
    ):
        add_step(doc, text, steps)
    add_callout(
        doc,
        "Tải lại",
        "Nút Tải lại lấy lại dữ liệu đã lưu từ máy chủ. Nếu đang có thay đổi chưa lưu, hãy kiểm tra kỹ trước khi dùng vì dữ liệu trên màn hình sẽ được nạp lại.",
        kind="warning",
    )

    add_heading(doc, "5.4. Ảnh hưởng của công tắc Chat nội bộ", 2)
    add_body(
        doc,
        "Công tắc Bật Chat nội bộ ở cấp doanh nghiệp là công tắc cha. Khi doanh nghiệp tắt Chat nội bộ, chi nhánh không thể bật lại. Trên thẻ Chi nhánh, dòng này sẽ bị vô hiệu hóa và hiển thị lý do.",
    )

    doc.add_page_break()
    add_heading(doc, "6. Cấu hình theo chi nhánh", 1)
    add_body(
        doc,
        "Cấu hình theo chi nhánh chỉ áp dụng cho chi nhánh được chọn. Hãy dùng cấp này khi chi nhánh có lịch làm việc, quy tắc vận hành hoặc chính sách khác với phần còn lại của doanh nghiệp.",
    )
    add_figure_pair(
        doc,
        "02-chi-nhanh-trai.png",
        "02-chi-nhanh-phai.png",
        "Hình 2. Thẻ Chi nhánh với chi nhánh Trụ sở Hà Nội đang được chọn (khung giao diện rộng 1920 px).",
        "Thẻ Chi nhánh của màn hình Cấu hình hệ thống VIN-HRM",
    )
    add_callout(
        doc,
        "Điểm khác với thẻ Doanh nghiệp",
        "Thẻ Chi nhánh có thêm ô chọn Chi nhánh ở đầu nội dung. Mọi thay đổi và thao tác lưu chỉ áp dụng cho chi nhánh đang hiển thị trong ô này.",
    )

    add_heading(doc, "6.1. Chọn đúng chi nhánh", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Chọn thẻ Chi nhánh.",
        "Mở danh sách Chi nhánh ở đầu trang.",
        "Chọn đúng tên chi nhánh cần cấu hình. Trong lần kiểm tra, phần mềm tự chọn chi nhánh đang hoạt động đầu tiên là Trụ sở Hà Nội.",
        "Chờ danh sách cấu hình tải xong rồi mới chỉnh sửa.",
    ):
        p = add_step(doc, text, steps)
        if text.startswith("Chọn đúng tên chi nhánh"):
            p.paragraph_format.page_break_before = True
    add_figure_pair(
        doc,
        "03-chon-chi-nhanh-trai.png",
        "03-chon-chi-nhanh-phai.png",
        "Hình 3. Danh sách chọn chi nhánh đang hoạt động (khung giao diện rộng 1920 px).",
        "Danh sách chọn chi nhánh trong màn hình Cấu hình hệ thống VIN-HRM",
    )
    add_callout(
        doc,
        "Kiểm tra tên chi nhánh",
        "Không dựa vào vị trí đầu tiên trong danh sách. Luôn đọc lại tên chi nhánh đang hiển thị trước khi nhấn Lưu thay đổi.",
        kind="warning",
    )

    add_heading(doc, "6.2. Hiểu giá trị ban đầu trên thẻ Chi nhánh", 2)
    add_body(
        doc,
        "Khi chi nhánh chưa có giá trị riêng, màn hình hiển thị giá trị hiệu lực từ doanh nghiệp; nếu doanh nghiệp cũng chưa cấu hình, màn hình hiển thị mặc định hệ thống. Vì vậy một giá trị hiển thị sẵn không đồng nghĩa với việc chi nhánh đã có bản ghi ghi đè.",
    )
    add_callout(
        doc,
        "Không có nhãn nguồn",
        "Giao diện hiện chưa chỉ rõ từng giá trị đến từ chi nhánh, doanh nghiệp hay mặc định. Khi cần xác định nguồn chính xác, hãy kết hợp thẻ Lịch sử và hỗ trợ kỹ thuật.",
    )

    add_heading(doc, "6.3. Chỉnh sửa và lưu", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Sau khi chọn chi nhánh, tìm cấu hình bằng ô tìm kiếm hoặc bộ lọc nhóm.",
        "Thay đổi đúng giá trị cần ghi đè. Các trường chỉ áp dụng ở cấp doanh nghiệp sẽ không xuất hiện.",
        "Kiểm tra lại tên chi nhánh, giá trị mới và ảnh hưởng nghiệp vụ.",
        "Chọn Lưu thay đổi.",
        "Khi thành công, hệ thống hiển thị 'Lưu cấu hình chi nhánh thành công'.",
        "Mở Lịch sử, chọn Theo chi nhánh và đúng chi nhánh để xác nhận.",
    ):
        p = add_step(doc, text, steps)
        if text.startswith("Khi thành công"):
            p.paragraph_format.keep_with_next = True

    add_heading(doc, "6.4. Ví dụ: đổi ngày xếp lịch cho một chi nhánh", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Chọn thẻ Chi nhánh và chọn Trụ sở Hà Nội.",
        "Lọc nhóm Lịch làm việc.",
        "Giữ Chu kỳ tự động xếp lịch làm việc là Hằng tuần.",
        "Chọn ngày phù hợp tại dòng Ngày tự động xếp lịch trong tuần.",
        "Chọn Lưu thay đổi và kiểm tra lịch sử theo chi nhánh.",
    ):
        add_step(doc, text, steps)

    add_callout(
        doc,
        "Khôi phục kế thừa",
        "Giao diện hiện không có nút 'Xóa ghi đè' hoặc 'Dùng lại giá trị doanh nghiệp'. Nhập một giá trị trùng với doanh nghiệp không đồng nghĩa với xóa ghi đè; chi nhánh có thể vẫn giữ một giá trị riêng. Nếu cần quay lại cơ chế kế thừa động, hãy liên hệ quản trị kỹ thuật.",
        kind="warning",
    )

    doc.add_page_break()
    add_heading(doc, "7. Kiểm tra lịch sử thay đổi", 1)
    add_body(
        doc,
        "Thẻ Lịch sử giúp kiểm tra một cấu hình đã thay đổi khi nào và thay đổi từ giá trị nào sang giá trị nào. Đây là bước bắt buộc sau các thay đổi có ảnh hưởng lớn.",
    )
    add_figure_pair(
        doc,
        "04-lich-su-trai.png",
        "04-lich-su-phai.png",
        "Hình 4. Lịch sử cấu hình theo doanh nghiệp (khung giao diện rộng 1920 px).",
        "Thẻ Lịch sử của màn hình Cấu hình hệ thống VIN-HRM",
    )

    add_heading(doc, "7.1. Xem lịch sử theo doanh nghiệp", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Chọn thẻ Lịch sử.",
        "Chọn Theo doanh nghiệp.",
        "Đọc các cột Thời điểm, Tên cấu hình, Giá trị cũ và Giá trị mới.",
        "Chọn Tải lại lịch sử nếu vừa lưu cấu hình nhưng chưa thấy bản ghi mới.",
    ):
        add_step(doc, text, steps)

    add_heading(doc, "7.2. Xem lịch sử theo chi nhánh", 2)
    steps = new_numbering_instance(doc, decimal_abstract_id)
    for text in (
        "Chọn Theo chi nhánh.",
        "Chọn chi nhánh cần kiểm tra trong danh sách xuất hiện.",
        "Đối chiếu tên cấu hình, giá trị cũ và giá trị mới.",
        "Chọn Xóa bộ lọc để quay về phạm vi Theo doanh nghiệp khi cần.",
    ):
        add_step(doc, text, steps)
    add_callout(
        doc,
        "Bảo vệ thông tin nhạy cảm",
        "Lịch sử che giá trị của cấu hình mật khẩu bằng dấu ******. Không chụp, sao chép hoặc gửi mật khẩu thật qua tài liệu hỗ trợ.",
        kind="success",
    )

    add_heading(doc, "8. Tình huống thường gặp và cách xử lý", 1)
    add_table(
        doc,
        ["Tình huống", "Nguyên nhân thường gặp", "Cách xử lý"],
        [
            ("Không thấy cấu hình cần tìm", "Đang có từ khóa hoặc bộ lọc nhóm.", "Chọn Xóa bộ lọc, sau đó tìm lại bằng tên hoặc mã khóa."),
            ("Không có dữ liệu chi nhánh", "Chưa chọn chi nhánh hoặc không có chi nhánh đang hoạt động.", "Chọn chi nhánh; nếu danh sách trống, kiểm tra trạng thái chi nhánh trong phần Tổ chức."),
            ("Không lưu được ngày trong tháng", "Giá trị ngoài khoảng 1 đến 31.", "Nhập số từ 1 đến 31 rồi lưu lại."),
            ("Không bật được Chat ở chi nhánh", "Chat đã tắt ở cấp doanh nghiệp.", "Đánh giá chính sách và bật ở doanh nghiệp trước; chi nhánh không được vượt công tắc cha."),
            ("Giá trị thay đổi sau khi bấm Tải lại", "Có chỉnh sửa chưa lưu hoặc dữ liệu máy chủ đã thay đổi.", "Nhập lại nếu cần, rà soát và lưu; kiểm tra lịch sử để xác định thay đổi gần nhất."),
            ("Không biết giá trị đến từ cấp nào", "Giao diện chỉ hiển thị giá trị hiệu lực.", "Đối chiếu doanh nghiệp, chi nhánh, lịch sử; liên hệ hỗ trợ kỹ thuật nếu cần xác định bản ghi ghi đè."),
            ("Muốn trả chi nhánh về kế thừa", "Không có nút xóa ghi đè trên giao diện.", "Không chỉ nhập lại giá trị giống doanh nghiệp. Gửi yêu cầu cho quản trị kỹ thuật để xóa ghi đè đúng cách."),
        ],
        [2400, 2900, 4060],
        font_size=9,
    )

    add_heading(doc, "9. Checklist vận hành an toàn", 1)
    add_heading(doc, "9.1. Trước khi lưu", 2)
    for text in (
        "Đang ở đúng thẻ Doanh nghiệp hoặc Chi nhánh.",
        "Nếu ở thẻ Chi nhánh, tên chi nhánh đã được kiểm tra lại.",
        "Đã hiểu ảnh hưởng của cấu hình đối với lịch làm việc, chấm công, lương, Chat, bảo hiểm hoặc OKR.",
        "Giá trị nhập đúng kiểu dữ liệu và đúng phạm vi.",
        "Đã thông báo cho bộ phận liên quan nếu thay đổi có tác động rộng.",
        "Không có thông tin bí mật đang hiển thị rõ trên màn hình chia sẻ.",
    ):
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "9.2. Sau khi lưu", 2)
    for text in (
        "Đã thấy thông báo lưu thành công.",
        "Đã tải lại hoặc mở lại thẻ để xác nhận giá trị hiệu lực.",
        "Đã kiểm tra lịch sử đúng phạm vi.",
        "Đã thử nghiệp vụ liên quan nếu thay đổi có rủi ro cao.",
        "Đã ghi nhận lý do thay đổi theo quy trình nội bộ của doanh nghiệp.",
    ):
        p = add_bullet(doc, text, bullet_num_id)
        if text.startswith("Đã thấy thông báo"):
            p.paragraph_format.page_break_before = True

    add_callout(
        doc,
        "Nguyên tắc",
        "Chỉ thay đổi những cấu hình cần thiết. Không chỉnh nhiều nhóm trong cùng một lần nếu không có kế hoạch kiểm tra, vì sẽ khó xác định nguyên nhân khi nghiệp vụ thay đổi.",
        kind="warning",
    )

    add_heading(doc, "10. Phụ lục", 1)
    add_heading(doc, "10.1. Thuật ngữ", 2)
    add_table(
        doc,
        ["Thuật ngữ", "Giải thích"],
        [
            ("Cấu hình mặc định", "Giá trị nền do hệ thống định nghĩa sẵn."),
            ("Ghi đè", "Giá trị ở cấp doanh nghiệp hoặc chi nhánh thay thế giá trị cấp thấp hơn trong thứ tự ưu tiên."),
            ("Giá trị hiệu lực", "Giá trị cuối cùng hệ thống sử dụng sau khi áp dụng thứ tự chi nhánh, doanh nghiệp, mặc định."),
            ("Mã khóa", "Mã nhận diện duy nhất của cấu hình, ví dụ System.DateFormat."),
            ("Cấu hình chỉ ở doanh nghiệp", "Cấu hình không được phép ghi đè theo chi nhánh."),
        ],
        [2400, 6960],
    )

    add_heading(doc, "10.2. Thông báo thành công cần ghi nhận", 2)
    for text in (
        "Lưu cấu hình doanh nghiệp thành công",
        "Lưu cấu hình chi nhánh thành công",
    ):
        add_bullet(doc, text, bullet_num_id)

    add_heading(doc, "10.3. Nguồn đối chiếu", 2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Tài liệu đặc tả nghiệp vụ: ")
    add_hyperlink(
        p,
        "02. Cấu hình v1.0",
        "https://docs.google.com/document/d/11dW_3cQNBK_E86RgZMj4v8VdcBwwzfNQYNPvenrjcPU/edit",
    )
    add_body(
        doc,
        "Giao diện và hành vi được kiểm tra từ mã nguồn VIN-HRM chạy cục bộ, kết nối môi trường dữ liệu được cấu hình trong dự án, vào ngày 13/08/2026.",
    )
    add_callout(
        doc,
        "Lưu ý phiên bản",
        "Tên nhóm, số lượng cấu hình và giá trị có thể thay đổi theo phiên bản phần mềm hoặc theo dữ liệu doanh nghiệp. Khi giao diện khác tài liệu, ưu tiên kiểm tra phiên bản đang triển khai và liên hệ bộ phận quản trị hệ thống.",
    )
    end_spacer = doc.add_paragraph()
    end_spacer.paragraph_format.space_before = Pt(0)
    end_spacer.paragraph_format.space_after = Pt(0)
    end_spacer.paragraph_format.line_spacing = Pt(1)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)

    doc.settings.update_fields_on_open = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
